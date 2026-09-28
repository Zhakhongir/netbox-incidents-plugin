"""
Статистика по инцидентам.

Строки — группы сайтов (sitegroup), внутри каждой группы отдельная строка на
каждый тип объекта (Стойка/Устройство/Модуль/Интерфейс, при наличии) плюс
строка "Итого по группе". 
Колонки — причины (cause) из IncidentCauseChoices, каждая с парой "Кол-во"/"Общее время"
"""
import datetime as dt
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from . import reports_common as rc
from .choices import IncidentCauseChoices

BOLD = Font(bold=True)
GRAY_FILL = PatternFill("solid", fgColor="D3D3D3")
TOTAL_FILL = PatternFill("solid", fgColor="B8CCE4")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

CAUSES = [c[0] for c in IncidentCauseChoices.CHOICES]
LABEL_COLS = 2  # A: Группа, B: Тип объекта / Итого
DATA_COLS = len(CAUSES) * 2
TOTAL_COLS = LABEL_COLS + DATA_COLS


def _apply_borders(ws, min_row, max_row):
    rc.apply_borders(ws, min_row, max_row, 1, TOTAL_COLS)


def _write_headers(ws):
    ws.merge_cells("A1:A2")
    a = ws["A1"]
    a.value = "Группа"
    a.font = BOLD
    a.alignment = CENTER

    ws.merge_cells("B1:B2")
    b = ws["B1"]
    b.value = "Тип объекта"
    b.font = BOLD
    b.alignment = CENTER

    col = LABEL_COLS + 1
    for cause in CAUSES:
        c_from = get_column_letter(col)
        c_to = get_column_letter(col + 1)
        ws.merge_cells(f"{c_from}1:{c_to}1")
        top = ws[f"{c_from}1"]
        top.value = cause
        top.font = BOLD
        top.alignment = CENTER

        cnt_cell = ws.cell(row=2, column=col, value="Кол-во")
        dur_cell = ws.cell(row=2, column=col + 1, value="Общее время")
        cnt_cell.font = BOLD
        dur_cell.font = BOLD
        cnt_cell.alignment = CENTER
        dur_cell.alignment = CENTER
        col += 2

    _apply_borders(ws, 1, 2)


def _write_data_row(ws, row, group_label, obj_label, counts, durations, bold=False, fill=None):
    gcell = ws.cell(row=row, column=1, value=group_label)
    ocell = ws.cell(row=row, column=2, value=obj_label)
    if bold:
        gcell.font = BOLD
        ocell.font = BOLD
    if fill:
        gcell.fill = fill
        ocell.fill = fill

    col = LABEL_COLS + 1
    for cause in CAUSES:
        cnt = counts.get(cause, 0)
        dur = durations.get(cause, dt.timedelta())
        c_cell = ws.cell(row=row, column=col, value=cnt)
        d_cell = ws.cell(row=row, column=col + 1, value=rc.format_duration(dur))
        if bold:
            c_cell.font = BOLD
            d_cell.font = BOLD
        if fill:
            c_cell.fill = fill
            d_cell.fill = fill
        col += 2

    _apply_borders(ws, row, row)


def generate(queryset, date_from: dt.date, date_to: dt.date):
    period_start, period_end = rc.make_period(date_from, date_to)
    resolved = rc.resolve_incidents(queryset, period_start, period_end)

    # группа сайтов -> тип объекта -> {counts: {cause: n}, durations: {cause: timedelta}}
    grouped: "OrderedDict" = OrderedDict()
    for item in resolved:
        by_obj = grouped.setdefault(item.group_name, OrderedDict())
        bucket = by_obj.setdefault(item.object_type, {'counts': {}, 'durations': {}})
        bucket['counts'][item.cause] = bucket['counts'].get(item.cause, 0) + 1
        bucket['durations'][item.cause] = bucket['durations'].get(item.cause, dt.timedelta()) + item.duration

    wb = Workbook()
    ws = wb.active
    ws.title = "Статистика"

    _write_headers(ws)
    row = 3

    grand_counts: dict = {}
    grand_durations: dict = {}

    for group_name, by_obj in grouped.items():
        group_start_row = row
        group_counts: dict = {}
        group_durations: dict = {}

        for obj_type in rc.OBJECT_TYPE_ORDER:
            bucket = by_obj.get(obj_type)
            if not bucket:
                continue

            _write_data_row(
                ws, row, "", rc.OBJECT_TYPE_LABELS[obj_type],
                bucket['counts'], bucket['durations'],
            )
            row += 1

            for cause, cnt in bucket['counts'].items():
                group_counts[cause] = group_counts.get(cause, 0) + cnt
                grand_counts[cause] = grand_counts.get(cause, 0) + cnt
            for cause, dur in bucket['durations'].items():
                group_durations[cause] = group_durations.get(cause, dt.timedelta()) + dur
                grand_durations[cause] = grand_durations.get(cause, dt.timedelta()) + dur

        _write_data_row(
            ws, row, "", f"Итого по группе «{group_name}»",
            group_counts, group_durations, bold=True, fill=TOTAL_FILL,
        )
        row += 1

        if row - 1 >= group_start_row:
            ws.merge_cells(start_row=group_start_row, start_column=1, end_row=row - 1, end_column=1)
            gcell = ws.cell(row=group_start_row, column=1)
            gcell.value = group_name
            gcell.font = BOLD
            gcell.alignment = CENTER

    _write_data_row(
        ws, row, "", "Итого",
        grand_counts, grand_durations, bold=True, fill=GRAY_FILL,
    )

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 26
    col = LABEL_COLS + 1
    for _ in CAUSES:
        ws.column_dimensions[get_column_letter(col)].width = 10
        ws.column_dimensions[get_column_letter(col + 1)].width = 15
        col += 2

    return rc.xlsx_response(
        wb,
        f"statistic_report_{date_from.isoformat()}_{date_to.isoformat()}.xlsx",
    )
