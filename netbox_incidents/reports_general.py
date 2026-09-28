"""Общий отчёт по инцидентам — на базе Incident (ORM)."""
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Font

from . import reports_common as rc

BOLD = Font(bold=True)

COMMON_HEADERS = [
    "Место установки", "Наименование оборудования",
    "Начало простоя", "Конец простоя", "Время простоя",
    "Причина простоя", "Кому сообщили", "Примечание",
]
AUTOSIZE_COLS = ["A", "B", "C", "D", "F", "H"]


def generate(queryset, date_from: dt.date, date_to: dt.date):
    period_start, period_end = rc.make_period(date_from, date_to)
    resolved = rc.resolve_incidents(queryset, period_start, period_end)

    groups: dict = {}
    for r in resolved:
        groups.setdefault(r.group_name, []).append(r)

    wb = Workbook()
    ws = wb.active
    ws.title = "Отчет"

    if date_from == date_to:
        title = f"Общий отчёт за {date_from.strftime('%d.%m.%Y')}"
    else:
        title = (
            f"Общий отчёт с {date_from.strftime('%d.%m.%Y')} "
            f"по {date_to.strftime('%d.%m.%Y')}"
        )
    ws.cell(row=1, column=1, value=title)

    last_col = len(COMMON_HEADERS)

    row = 3
    for col_i, label in enumerate(COMMON_HEADERS, start=1):
        ws.cell(row=row, column=col_i, value=label).font = BOLD
    rc.apply_borders(ws, row, row, 1, last_col)
    row += 1

    for group, items in groups.items():
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
        cell = ws.cell(row=row, column=1)
        cell.value = group
        cell.font = BOLD
        rc.apply_borders(ws, row, row, 1, last_col)
        row += 1

        for item in items:
            # Комментарий весь текст целиком идёт в "Примечание".
            duration_str = rc.format_duration(item.duration) if item.end else ""

            ws.cell(row=row, column=1, value=item.site_name)
            ws.cell(row=row, column=2, value=item.equipment_name)
            ws.cell(row=row, column=3, value=rc.format_date(item.start))
            ws.cell(row=row, column=4, value=rc.format_date(item.end))
            ws.cell(row=row, column=5, value=duration_str)
            ws.cell(row=row, column=6, value=item.cause)
            ws.cell(row=row, column=7, value=item.contact_name)
            ws.cell(row=row, column=8, value=item.comments)
            rc.apply_borders(ws, row, row, 1, last_col)
            row += 1

    rc.autosize_columns(ws, AUTOSIZE_COLS)

    return rc.xlsx_response(
        wb,
        f"general_report_{date_from.isoformat()}_{date_to.isoformat()}.xlsx",
    )
