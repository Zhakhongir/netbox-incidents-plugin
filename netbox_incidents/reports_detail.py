"""
Детальный отчёт по инцидентам.

Группировка (3 уровня): группа сайтов (sitegroup) → причина (cause) → тип объекта (Стойка/Устройство/Модуль/Интерфейс)
"""
import datetime as dt
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from . import reports_common as rc
from .choices import IncidentCauseChoices

BOLD = Font(bold=True)
GRAY_FILL = PatternFill("solid", fgColor="D3D3D3")
CAUSE_FILL = PatternFill("solid", fgColor="DCE6F1")
SUBGROUP_FILL = PatternFill("solid", fgColor="F2F2F2")
OBJTYPE_TOTAL_FILL = PatternFill("solid", fgColor="E2EFDA")
TOTAL_FILL = PatternFill("solid", fgColor="B8CCE4")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

DETAIL_HEADERS = [
    "Место установки", "Наименование оборудования", "Интерфейс", "IP адрес",
    "Начало простоя", "Конец простоя", "Время простоя", "Примечание",
]
DETAIL_COLS = len(DETAIL_HEADERS)
LAST_COL_LETTER = get_column_letter(DETAIL_COLS)

CAUSE_ORDER = [c[0] for c in IncidentCauseChoices.CHOICES]


def _cause_sort_key(cause_value):
    try:
        return CAUSE_ORDER.index(cause_value)
    except ValueError:
        return len(CAUSE_ORDER)


def _write_header_row(ws, row):
    for col_i, label in enumerate(DETAIL_HEADERS, start=1):
        cell = ws.cell(row=row, column=col_i, value=label)
        cell.font = BOLD
        cell.fill = GRAY_FILL
        cell.alignment = CENTER
    rc.apply_borders(ws, row, row, 1, DETAIL_COLS)


def _write_incident_row(ws, row, item):
    ws.cell(row=row, column=1, value=item.site_name)
    ws.cell(row=row, column=2, value=item.equipment_name)
    ws.cell(row=row, column=3, value=item.interface_name)
    ws.cell(row=row, column=4, value=item.ip_address)
    ws.cell(row=row, column=5, value=rc.format_date(item.start))
    ws.cell(row=row, column=6, value=rc.format_date(item.end) if item.end else "открытый")
    ws.cell(row=row, column=7, value=rc.format_duration(item.duration))
    desc_cell = ws.cell(row=row, column=8, value=item.comments)
    desc_cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 28
    rc.apply_borders(ws, row, row, 1, DETAIL_COLS)


def generate(queryset, date_from: dt.date, date_to: dt.date):
    period_start, period_end = rc.make_period(date_from, date_to)
    resolved = rc.resolve_incidents(queryset, period_start, period_end)

    # группа сайтов -> причина -> тип объекта -> [инциденты]
    grouped: "OrderedDict" = OrderedDict()
    for item in resolved:
        by_cause = grouped.setdefault(item.group_name, OrderedDict())
        by_obj_type = by_cause.setdefault(item.cause, OrderedDict())
        by_obj_type.setdefault(item.object_type, []).append(item)

    wb = Workbook()
    ws = wb.active
    ws.title = "Детальный отчёт"

    row = 1
    _write_header_row(ws, row)
    row += 1

    for group_name, by_cause in grouped.items():
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=DETAIL_COLS)
        gcell = ws.cell(row=row, column=1)
        gcell.value = group_name
        gcell.font = Font(bold=True, size=12)
        gcell.fill = GRAY_FILL
        gcell.alignment = CENTER
        rc.apply_borders(ws, row, row, 1, DETAIL_COLS)
        row += 1

        group_total_duration = dt.timedelta()
        ordered_causes = sorted(by_cause.keys(), key=_cause_sort_key)

        for cause in ordered_causes:
            by_obj_type = by_cause[cause]
            cause_items = [i for items in by_obj_type.values() for i in items]
            cause_duration = sum((i.duration for i in cause_items), dt.timedelta())
            cause_count = len(cause_items)

            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=DETAIL_COLS)
            ccell = ws.cell(row=row, column=1)
            ccell.value = f"{cause}   [{cause_count} сл. / {rc.format_duration(cause_duration)}]"
            ccell.font = BOLD
            ccell.fill = CAUSE_FILL
            ccell.alignment = Alignment(horizontal="left", vertical="center")
            rc.apply_borders(ws, row, row, 1, DETAIL_COLS)
            row += 1

            for obj_type in rc.OBJECT_TYPE_ORDER:
                items = by_obj_type.get(obj_type)
                if not items:
                    continue

                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=DETAIL_COLS)
                scell = ws.cell(row=row, column=1)
                scell.value = rc.OBJECT_TYPE_LABELS[obj_type]
                scell.font = Font(bold=True, italic=True, size=10)
                scell.fill = SUBGROUP_FILL
                scell.alignment = Alignment(horizontal="left", vertical="center")
                rc.apply_borders(ws, row, row, 1, DETAIL_COLS)
                row += 1

                _write_header_row(ws, row)
                row += 1

                for item in items:
                    _write_incident_row(ws, row, item)
                    row += 1

                obj_duration = sum((i.duration for i in items), dt.timedelta())
                obj_count = len(items)

                obj_subtotal_row = row
                ws.merge_cells(start_row=obj_subtotal_row, start_column=1, end_row=obj_subtotal_row, end_column=6)
                olabel = ws.cell(row=obj_subtotal_row, column=1)
                olabel.value = f"Итого по «{rc.OBJECT_TYPE_LABELS[obj_type]}»   [{obj_count} сл.]"
                olabel.font = BOLD
                olabel.fill = OBJTYPE_TOTAL_FILL
                olabel.alignment = Alignment(horizontal="right")
                odur = ws.cell(row=obj_subtotal_row, column=7, value=rc.format_duration(obj_duration))
                odur.font = BOLD
                odur.fill = OBJTYPE_TOTAL_FILL
                ws.cell(row=obj_subtotal_row, column=8).fill = OBJTYPE_TOTAL_FILL
                rc.apply_borders(ws, obj_subtotal_row, obj_subtotal_row, 1, DETAIL_COLS)
                row += 1

            subtotal_row = row
            ws.merge_cells(start_row=subtotal_row, start_column=1, end_row=subtotal_row, end_column=6)
            slabel = ws.cell(row=subtotal_row, column=1)
            slabel.value = f"Итого по «{cause}»"
            slabel.font = BOLD
            slabel.fill = TOTAL_FILL
            slabel.alignment = Alignment(horizontal="right")
            sdur = ws.cell(row=subtotal_row, column=7, value=rc.format_duration(cause_duration))
            sdur.font = BOLD
            sdur.fill = TOTAL_FILL
            ws.cell(row=subtotal_row, column=8).fill = TOTAL_FILL
            rc.apply_borders(ws, subtotal_row, subtotal_row, 1, DETAIL_COLS)
            row += 1

            group_total_duration += cause_duration

        total_row = row
        ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=6)
        tlabel = ws.cell(row=total_row, column=1)
        tlabel.value = f"Итого по группе «{group_name}»"
        tlabel.font = Font(bold=True)
        tlabel.fill = GRAY_FILL
        tlabel.alignment = Alignment(horizontal="right")
        tdur = ws.cell(row=total_row, column=7, value=rc.format_duration(group_total_duration))
        tdur.font = Font(bold=True)
        tdur.fill = GRAY_FILL
        ws.cell(row=total_row, column=8).fill = GRAY_FILL
        rc.apply_borders(ws, total_row, total_row, 1, DETAIL_COLS)
        row += 2

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 80

    return rc.xlsx_response(
        wb,
        f"detail_report_{date_from.isoformat()}_{date_to.isoformat()}.xlsx",
    )
